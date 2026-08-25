"""The runner/failure research database. Every candidate the scanner ever
flags gets a row here - whether the bot actually traded it or not - plus
every trade's full outcome (exit reason, P&L, how far it ran for and
against you). This is what phase 4 (replay/slippage) and phase 5 (paper
sample analysis) actually query against; the text log is for watching the
bot live, this is for research after the fact.

Plain sqlite3 (stdlib, no extra dependency) - one file, no server, trivial
to open with any SQL tool or pandas.read_sql for analysis.
"""

import json
import sqlite3
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

DEFAULT_HISTORICAL_WORKBOOK = (
    Path(__file__).resolve().parent.parent
    / "research"
    / "MR_Michael_Historical_Runners_Failures_Controls.xlsx"
)

DB_PATH = Path(__file__).resolve().parent.parent / "state" / "tradeboss.db"
ET = ZoneInfo("America/New_York")

_SCHEMA = """
CREATE TABLE IF NOT EXISTS candidate_snapshots (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL,
    symbol TEXT NOT NULL,
    last_price REAL,
    prev_close REAL,
    gain_pct REAL,
    relative_volume REAL,
    float_shares REAL,
    float_source TEXT,
    float_stale INTEGER,
    float_stale_reasons TEXT,
    rotations_since_open REAL,
    recent_turnover_rate REAL,
    float_adjusted_relative_volume REAL,
    buzz_messages_recent INTEGER,
    trending INTEGER,
    market_cap REAL,
    spread_pct REAL,
    bid_ask_imbalance REAL,
    rsi REAL,
    score REAL,
    entered INTEGER
);

CREATE TABLE IF NOT EXISTS halt_sightings (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    ts TEXT NOT NULL,
    symbol TEXT NOT NULL,
    last_known_price REAL
);

CREATE TABLE IF NOT EXISTS trades (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    symbol TEXT NOT NULL,
    entry_ts TEXT NOT NULL,
    entry_price REAL,
    qty REAL,
    exit_ts TEXT,
    exit_price REAL,
    exit_reason TEXT,
    exit_category TEXT,
    pnl REAL,
    peak_price REAL,
    trough_price REAL,
    mfe_pct REAL,
    mae_pct REAL,
    dry_run INTEGER,
    entry_score REAL,
    entry_float_shares REAL,
    entry_rotations_since_open REAL,
    entry_buzz_messages_recent REAL,
    entry_market_cap REAL,
    entry_rsi REAL
);

CREATE INDEX IF NOT EXISTS idx_snapshots_symbol_ts ON candidate_snapshots(symbol, ts);
CREATE INDEX IF NOT EXISTS idx_trades_symbol ON trades(symbol);

-- Reference tables imported from research/MR_Michael_Historical_Runners_
-- Failures_Controls.xlsx (see import_historical_workbook below). These are
-- hand-researched, cited historical cases - not written to by the live
-- bot - kept here so they're queryable/joinable alongside the bot's own
-- candidate_snapshots/trades once real point-in-time features accumulate.
-- The workbook's own no-lookahead discipline applies to any future code
-- that joins against these: never match on outcome fields, and note that
-- Feature_Snapshots/Outcomes in the source workbook are schemas only -
-- no point-in-time data has actually been ingested yet.
CREATE TABLE IF NOT EXISTS historical_events (
    event_id TEXT PRIMARY KEY,
    ticker TEXT,
    issuer TEXT,
    event_date TEXT,
    event_window TEXT,
    exchange TEXT,
    label TEXT,
    taxonomy TEXT,
    catalyst_type TEXT,
    catalyst_quality TEXT,
    headline_move_pct REAL,
    volume_shares REAL,
    reference_price REAL,
    reverse_split_state TEXT,
    foreign_issuer TEXT,
    regulatory_outcome TEXT,
    failure_mode TEXT,
    status TEXT,
    source_1 TEXT,
    source_2 TEXT,
    notes TEXT
);

CREATE TABLE IF NOT EXISTS historical_matched_controls (
    match_id TEXT PRIMARY KEY,
    treated_event_id TEXT,
    control_event_id TEXT,
    match_class TEXT,
    match_status TEXT,
    rationale TEXT,
    hard_match_dimensions TEXT,
    missing_for_statistical_match TEXT,
    match_score REAL
);

CREATE TABLE IF NOT EXISTS historical_match_rules (
    variable TEXT,
    transform_bucket TEXT,
    hard_caliper TEXT,
    default_weight TEXT,
    leakage_rule TEXT
);

CREATE TABLE IF NOT EXISTS historical_sources (
    source_id TEXT PRIMARY KEY,
    domain TEXT,
    source_type TEXT,
    url_or_reference TEXT,
    purpose TEXT
);
"""


def _connect() -> sqlite3.Connection:
    DB_PATH.parent.mkdir(exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.executescript(_SCHEMA)
    return conn


def _now_iso() -> str:
    return datetime.now(ET).isoformat()


def log_candidate_snapshot(candidate: dict, entered: bool, log):
    """Record a scanner sighting - called for every candidate that passed
    the filters in get_runner_candidates, regardless of whether the bot
    actually entered it (risk limits may have blocked it). This is the
    "runner" side of the runner/failure database."""
    try:
        with _connect() as conn:
            conn.execute(
                """INSERT INTO candidate_snapshots
                   (ts, symbol, last_price, prev_close, gain_pct, relative_volume,
                    float_shares, float_source, float_stale, float_stale_reasons,
                    rotations_since_open, recent_turnover_rate,
                    float_adjusted_relative_volume, buzz_messages_recent,
                    trending, market_cap, spread_pct, bid_ask_imbalance, rsi,
                    score, entered)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    _now_iso(),
                    candidate["symbol"],
                    candidate.get("last_price"),
                    candidate.get("prev_close"),
                    candidate.get("gain_pct"),
                    candidate.get("relative_volume"),
                    candidate.get("float_shares"),
                    candidate.get("float_source"),
                    int(bool(candidate.get("float_stale"))),
                    json.dumps(candidate.get("float_stale_reasons") or []),
                    candidate.get("rotations_since_open"),
                    candidate.get("recent_turnover_rate"),
                    candidate.get("float_adjusted_relative_volume"),
                    candidate.get("buzz_messages_recent"),
                    int(bool(candidate.get("trending"))),
                    candidate.get("market_cap"),
                    candidate.get("spread_pct"),
                    candidate.get("bid_ask_imbalance"),
                    candidate.get("rsi"),
                    candidate.get("score"),
                    int(entered),
                ),
            )
    except Exception as e:
        log.warning("Failed to log candidate snapshot for %s: %s", candidate.get("symbol"), e)


def log_halt_sighting(symbol: str, last_known_price: float | None, log):
    """Record that a symbol was seen halted - the "failure mode" data a
    resting stop order can't help with (see executor.py)."""
    try:
        with _connect() as conn:
            conn.execute(
                "INSERT INTO halt_sightings (ts, symbol, last_known_price) VALUES (?,?,?)",
                (_now_iso(), symbol, last_known_price),
            )
    except Exception as e:
        log.warning("Failed to log halt sighting for %s: %s", symbol, e)


def _exit_category(reason: str) -> str:
    reason = (reason or "").lower()
    if reason.startswith("stop loss"):
        return "stop_loss"
    if reason.startswith("take-profit cap"):
        return "take_profit_cap"
    if reason.startswith("trailing exit"):
        return "trailing"
    if reason.startswith("distribution"):
        return "distribution"
    if reason.startswith("force exit"):
        return "force_exit_time"
    return "other"


def log_trade(
    symbol: str,
    entry_ts: str,
    entry_price: float,
    qty: float,
    exit_price: float,
    exit_reason: str,
    pnl: float,
    peak_price: float,
    trough_price: float,
    dry_run: bool,
    entry_meta: dict,
    log,
):
    """Record a completed trade with its full outcome - the "failure" side
    of the database. `entry_meta` is whatever was captured about the
    candidate at entry time (score, float, rotations, buzz) so later
    analysis can ask "what made a winner different from a loser" without
    having to rejoin against the candidate_snapshots table."""
    mfe_pct = (peak_price - entry_price) / entry_price if entry_price else None
    mae_pct = (trough_price - entry_price) / entry_price if entry_price else None
    try:
        with _connect() as conn:
            conn.execute(
                """INSERT INTO trades
                   (symbol, entry_ts, entry_price, qty, exit_ts, exit_price,
                    exit_reason, exit_category, pnl, peak_price, trough_price,
                    mfe_pct, mae_pct, dry_run, entry_score, entry_float_shares,
                    entry_rotations_since_open, entry_buzz_messages_recent,
                    entry_market_cap, entry_rsi)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    symbol,
                    entry_ts,
                    entry_price,
                    qty,
                    _now_iso(),
                    exit_price,
                    exit_reason,
                    _exit_category(exit_reason),
                    pnl,
                    peak_price,
                    trough_price,
                    mfe_pct,
                    mae_pct,
                    int(dry_run),
                    entry_meta.get("score"),
                    entry_meta.get("float_shares"),
                    entry_meta.get("rotations_since_open"),
                    entry_meta.get("buzz_messages_recent"),
                    entry_meta.get("market_cap"),
                    entry_meta.get("rsi"),
                ),
            )
    except Exception as e:
        log.warning("Failed to log trade for %s: %s", symbol, e)


def _sheet_rows(ws) -> list[dict]:
    """Every research sheet in the workbook follows the same layout: title
    row, description row, blank row, header row, then data - so header is
    always row 4. Returns one dict per non-empty data row, keyed by
    header, skipping rows that are entirely blank (template rows with no
    data yet, e.g. Feature_Snapshots/Outcomes in the source workbook)."""
    header = [c.value for c in ws[4]]
    rows = []
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(header, row)))
    return rows


def import_historical_workbook(path: Path = DEFAULT_HISTORICAL_WORKBOOK, log=None) -> dict:
    """Imports Event_Master, Matched_Controls, Match_Rules, and Sources
    from the research workbook into historical_* reference tables.
    Idempotent - re-running after the workbook is edited replaces prior
    rows rather than duplicating them. Returns a dict of table -> row
    count imported. Requires openpyxl (only needed for this one-time/
    occasional import, not for running the bot itself)."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    counts = {}

    with _connect() as conn:
        events = _sheet_rows(wb["Event_Master"])
        conn.execute("DELETE FROM historical_events")
        for r in events:
            conn.execute(
                """INSERT INTO historical_events
                   (event_id, ticker, issuer, event_date, event_window, exchange,
                    label, taxonomy, catalyst_type, catalyst_quality,
                    headline_move_pct, volume_shares, reference_price,
                    reverse_split_state, foreign_issuer, regulatory_outcome,
                    failure_mode, status, source_1, source_2, notes)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    r.get("event_id"), r.get("ticker"), r.get("issuer"),
                    str(r.get("event_date")) if r.get("event_date") is not None else None,
                    r.get("event_window"), r.get("exchange"), r.get("label"),
                    r.get("taxonomy"), r.get("catalyst_type"), r.get("catalyst_quality"),
                    r.get("headline_move_pct"), r.get("volume_shares"), r.get("reference_price"),
                    r.get("reverse_split_state"), r.get("foreign_issuer"),
                    r.get("regulatory_outcome"), r.get("failure_mode"), r.get("status"),
                    r.get("source_1"), r.get("source_2"), r.get("notes"),
                ),
            )
        counts["historical_events"] = len(events)

        controls = _sheet_rows(wb["Matched_Controls"])
        conn.execute("DELETE FROM historical_matched_controls")
        for r in controls:
            conn.execute(
                """INSERT INTO historical_matched_controls
                   (match_id, treated_event_id, control_event_id, match_class,
                    match_status, rationale, hard_match_dimensions,
                    missing_for_statistical_match, match_score)
                   VALUES (?,?,?,?,?,?,?,?,?)""",
                (
                    r.get("match_id"), r.get("treated_event_id"), r.get("control_event_id"),
                    r.get("match_class"), r.get("match_status"), r.get("rationale"),
                    r.get("hard_match_dimensions"), r.get("missing_for_statistical_match"),
                    r.get("match_score"),
                ),
            )
        counts["historical_matched_controls"] = len(controls)

        rules = _sheet_rows(wb["Match_Rules"])
        conn.execute("DELETE FROM historical_match_rules")
        for r in rules:
            conn.execute(
                """INSERT INTO historical_match_rules
                   (variable, transform_bucket, hard_caliper, default_weight, leakage_rule)
                   VALUES (?,?,?,?,?)""",
                (
                    r.get("Variable"), r.get("Transform / bucket"), r.get("Hard caliper"),
                    str(r.get("Default weight")), r.get("Leakage rule"),
                ),
            )
        counts["historical_match_rules"] = len(rules)

        sources = _sheet_rows(wb["Sources"])
        conn.execute("DELETE FROM historical_sources")
        for r in sources:
            conn.execute(
                """INSERT INTO historical_sources
                   (source_id, domain, source_type, url_or_reference, purpose)
                   VALUES (?,?,?,?,?)""",
                (
                    r.get("source_id"), r.get("domain"), r.get("source_type"),
                    r.get("url_or_reference"), r.get("purpose"),
                ),
            )
        counts["historical_sources"] = len(sources)

    if log:
        log.info("Imported historical workbook: %s", counts)
    return counts
